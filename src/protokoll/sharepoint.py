"""Genererar SharePoint-/Copilot-färdiga Word-dokument (ett per sammanträde)
ur ärendeindexet. Word (.docx) indexeras tillförlitligt av Microsoft 365
Copilot, och den strukturerade datan (§, rubrik, diarienummer, beslut) ger
rena, sökvänliga dokument – betydligt bättre grounding än råa PDF:er.

Kör:  protokoll sharepoint
Skriver en mapp sharepoint/ med en .docx per protokoll, redo att laddas upp
till ett SharePoint-dokumentbibliotek.
"""

import json
from collections import defaultdict

from docx import Document

from . import config

OUT_DIR = config.ROOT / "sharepoint"


def _load() -> dict[str, list[dict]]:
    if not config.ARENDEN_JSONL.exists():
        raise RuntimeError("Inget index. Kör 'protokoll index' först.")
    per_protokoll: dict[str, list[dict]] = defaultdict(list)
    with config.ARENDEN_JSONL.open(encoding="utf-8") as fh:
        for line in fh:
            r = json.loads(line)
            per_protokoll[r["fil"]].append(r)
    return per_protokoll


def _build_doc(arenden: list[dict]) -> Document:
    arenden = sorted(arenden, key=lambda a: a["paragraf"])
    first = arenden[0]
    doc = Document()
    doc.add_heading(f"{first['namnd']} – sammanträde {first['datum']}", level=0)
    if first.get("kalla"):
        p = doc.add_paragraph()
        p.add_run("Källa: ").bold = True
        p.add_run(first["kalla"])
    doc.add_paragraph(
        "Detta dokument är en maskinell sammanställning av protokollets ärenden. "
        "Officiell handling: se källlänken ovan."
    ).italic = True

    for a in arenden:
        doc.add_heading(f"§ {a['paragraf']} – {a['rubrik']}", level=1)
        if a.get("dnr"):
            p = doc.add_paragraph()
            p.add_run("Diarienummer: ").bold = True
            p.add_run(a["dnr"])
        if a.get("beslut"):
            doc.add_heading("Beslut", level=2)
            doc.add_paragraph(a["beslut"])
        if a.get("text"):
            doc.add_heading("Ärendetext", level=2)
            for stycke in a["text"].split("\n\n"):
                stycke = _stada(stycke)
                if stycke:
                    doc.add_paragraph(stycke)
    return doc


def _stada(stycke: str) -> str:
    """Tar bort kvarvarande markdown-markörer (#, *) och tabellavgränsare."""
    rader = []
    for rad in stycke.splitlines():
        rad = rad.strip()
        if not rad or set(rad) <= set("|-: "):  # tom eller tabellavgränsare
            continue
        rad = rad.strip("#").strip().replace("**", "").strip()
        rader.append(rad)
    return " ".join(rader).strip()


def run() -> int:
    per_protokoll = _load()
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    n = 0
    for fil, arenden in per_protokoll.items():
        doc = _build_doc(arenden)
        namn = fil.rsplit(".", 1)[0] + ".docx"
        doc.save(str(OUT_DIR / namn))
        n += 1
    print(f"Skrev {n} Word-dokument till {OUT_DIR.relative_to(config.ROOT)}/")
    print("Ladda upp mappen till ett SharePoint-dokumentbibliotek och koppla det")
    print("som Knowledge i din Copilot-agent.")
    return 0
